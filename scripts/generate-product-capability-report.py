from pathlib import Path
from datetime import datetime

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
OUT_DIR = ROOT / "ICBC-WMP" / "报告"
OUT_PATH = OUT_DIR / "工行财富管理产品能力对标报告.docx"


NAVY = "0B2545"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
GRAY = "555555"
MUTED = "6B7280"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
CALLOUT = "F4F6F9"
GREEN = "1F6B4F"
GOLD = "7A5A00"
RED = "9B1C1C"
WHITE = "FFFFFF"


def set_font(run, name="Microsoft YaHei", size=None, color=None, bold=None, italic=None):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_style_font(style, name="Microsoft YaHei", size=11, color=None, bold=None):
    style.font.name = name
    style._element.rPr.rFonts.set(qn("w:ascii"), name)
    style._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    style.font.size = Pt(size)
    if color:
        style.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        style.font.bold = bold


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_borders(cell, color="D1D5DB", size=4):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        el = borders.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(size))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)


def set_cell_width(cell, width_dxa):
    cell.width = Inches(width_dxa / 1440)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths, indent=120):
    total = sum(widths)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.insert(0, tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            set_cell_width(cell, widths[min(idx, len(widths) - 1)])
            set_cell_margins(cell)
            set_cell_borders(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def mark_header_row(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant = OxmlElement("w:cantSplit")
    tr_pr.append(cant)


def clear_cell(cell):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.08
    return p


def write_cell(cell, text, bold=False, color=GRAY, size=9.2, align=WD_ALIGN_PARAGRAPH.LEFT):
    p = clear_cell(cell)
    p.alignment = align
    run = p.add_run(str(text))
    set_font(run, size=size, color=color, bold=bold)
    return p


def add_field(paragraph, instruction):
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr)
    run._r.append(fld_char2)
    set_font(run, size=9, color=MUTED)


def add_para(doc, text="", size=11, color=GRAY, bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.LEFT, before=0, after=6, line=1.10, style=None):
    p = doc.add_paragraph(style=style)
    p.alignment = align
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    if text:
        run = p.add_run(text)
        set_font(run, size=size, color=color, bold=bold, italic=italic)
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.5 + level * 0.25)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.167
    run = p.add_run(text)
    set_font(run, size=10.6, color=GRAY)
    return p


def add_heading(doc, text, level=1):
    style = f"Heading {level}"
    p = doc.add_paragraph(style=style)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    return p


def add_callout(doc, label, text, fill=CALLOUT, accent=BLUE):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360], indent=120)
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:space"), "0")
    left.set(qn("w:color"), accent)
    borders.append(left)
    p = clear_cell(cell)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    r1 = p.add_run(label + "  ")
    set_font(r1, size=10.5, color=accent, bold=True)
    r2 = p.add_run(text)
    set_font(r2, size=10.5, color=GRAY)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_matrix_table(doc, headers, rows, widths, header_fill=LIGHT_BLUE, font_size=8.8):
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_geometry(table, widths, indent=120)
    header = table.rows[0]
    mark_header_row(header)
    for i, label in enumerate(headers):
        set_cell_shading(header.cells[i], header_fill)
        write_cell(header.cells[i], label, bold=True, color=NAVY, size=font_size, align=WD_ALIGN_PARAGRAPH.CENTER)
    for row_data in rows:
        row = table.add_row()
        prevent_row_split(row)
        for i, value in enumerate(row_data):
            color = GRAY
            if i == len(row_data) - 1:
                if "高" in str(value):
                    color = GREEN
                elif "低" in str(value) or "未证实" in str(value) or "不在" in str(value):
                    color = RED
                elif "中" in str(value):
                    color = GOLD
            write_cell(row.cells[i], value, bold=(i == 0), color=color, size=font_size)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_source_list(doc, sources):
    for label, path, note in sources:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Inches(0.5)
        p.paragraph_format.first_line_indent = Inches(-0.25)
        p.paragraph_format.space_after = Pt(3)
        r1 = p.add_run(label + "：")
        set_font(r1, size=9.4, color=NAVY, bold=True)
        r2 = p.add_run(path)
        set_font(r2, name="Consolas", size=8.8, color=GRAY)
        if note:
            r3 = p.add_run("（" + note + "）")
            set_font(r3, size=9.2, color=MUTED)


def configure_styles(doc):
    normal = doc.styles["Normal"]
    set_style_font(normal, size=11, color=GRAY)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10
    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ]:
        style = doc.styles[name]
        set_style_font(style, size=size, color=color, bold=True)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.10
        style.paragraph_format.keep_with_next = True
    for name in ("List Bullet", "List Number"):
        style = doc.styles[name]
        set_style_font(style, size=10.6, color=GRAY)


def configure_section(section):
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hp.paragraph_format.space_after = Pt(0)
    r = hp.add_run("工行财富管理产品能力对标报告")
    set_font(r, size=8.8, color=MUTED, bold=True)
    r = hp.add_run("    研究快报")
    set_font(r, size=8.8, color=MUTED)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fp.paragraph_format.space_before = Pt(0)
    fp.paragraph_format.space_after = Pt(0)
    r = fp.add_run("ICBC-WMP  |  ")
    set_font(r, size=8.8, color=MUTED)
    add_field(fp, "PAGE")


def add_title_block(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("研究报告")
    set_font(r, size=10, color=BLUE, bold=True)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run("工商银行财富管理产品能力对标报告")
    set_font(r, size=24, color=NAVY, bold=True)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(13)
    r = p.add_run("对比图片标杆服务类型与 2026-08-21 工行公开理财产品爬取结果")
    set_font(r, size=13.2, color=GRAY)
    metadata = [
        ("研究对象", "工行个人网银公开理财产品列表（11/ALL，areaCode=0000）"),
        ("数据快照", "2026-08-21 14:48:10 UTC；138 页；1,375 条唯一产品代码"),
        ("对标材料", "用户提供的 4 张图片，归并为 13 类财富管理能力"),
        ("报告结论", "产品货架以现金管理和固定收益为核心；专业服务能力需另行接入数据源核验"),
    ]
    for label, value in metadata:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.05
        r = p.add_run(label + "：")
        set_font(r, size=9.8, color=NAVY, bold=True)
        r = p.add_run(value)
        set_font(r, size=9.8, color=GRAY)


def build_report():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_styles(doc)
    configure_section(doc.sections[0])

    core = (
        "基于本次抓取结果，工行个人网银理财货架已能较完整支撑现金管理、固定收益、部分混合/权益增强、"
        "FOF、ESG 以及有限的 QDII/外币理财筛选与比较；但图片中的保险规划、信托/遗产、税务、财富结构、"
        "机构客户和市场研究等，属于服务能力或其他产品域，不能由本次理财列表直接证明。"
    )
    add_title_block(doc)
    add_callout(doc, "核心结论", core, fill=CALLOUT, accent=BLUE)

    add_heading(doc, "一、管理层摘要", 1)
    add_para(doc, "本报告的判断对象不是“工行全部产品”，而是工行个人网银公开理财列表这一可复核的产品货架。"
             "因此，报告把“产品记录存在”“可由当前字段完成的动作”和“需要专业服务/其他数据源的能力”分开评价。")
    summary_rows = [
        ("产品供给", "较高", "现金管理、固定收益有大量记录，期限、起购金额、销售状态、风险标签等核心字段齐全。"),
        ("策略广度", "中等", "混合、FOF、ESG、红利/指数增强、少量指数挂钩和对冲策略可见；直接权益、另类和复杂衍生品证据弱。"),
        ("风险覆盖", "中等偏稳健", "1,375 条记录仅出现 PR1、PR2、PR3，无 PR4/PR5；不能据此证明高风险产品或直接交易能力。"),
        ("跨境暴露", "有限可用", "可见 QDII、美元、港币、欧元、澳元及全球配置名称，但数量和产品形态有限。"),
        ("综合财富服务", "未证实/需接入", "保险、信托、传承、税务、机构服务和市场研究不在本次清单字段范围内，不能用空关键词作否定结论。"),
    ]
    add_matrix_table(doc, ["维度", "完善度", "判断"], summary_rows, [1800, 1700, 5860], font_size=9.1)

    add_heading(doc, "二、对标范围与数据边界", 1)
    add_para(doc, "图片中的标签被作为能力对标框架使用，而不是作为工行已提供能力的证明。四张图片共展示 14 个标签；"
             "将“信托服务”并入“信托和遗产规划服务”后，形成 13 类能力：股权投资、固定收益策略、衍生产品交易、"
             "全权委托投资组合、另类投资、保险规划、信托和遗产规划服务、市场研究和分析、机构客户服务、全面传承规划、"
             "人寿保险、税务责任规划、财富结构设计。")
    add_callout(doc, "重要口径", "“本次清单未检出”只表示在公开理财列表中没有可核对记录，不等于工商银行全行不存在相关业务。"
                "对于服务型能力，建议标记为“待接入/需人工核验”，而不是自动推荐。", fill="FFF8E8", accent=GOLD)
    add_para(doc, "数据来源是工行个人网银公开列表入口及其分页列表服务；爬虫未自动打开每个产品的说明书、风险揭示书或交易页面。"
             "因此，名称关键词只能支持“产品货架上出现过该类标签”，不能替代产品主数据、合同条款、持仓、费率或交易资格核验。")

    add_heading(doc, "三、爬取结果画像", 1)
    profile_rows = [
        ("数据集", "ICBC-WMP/data/icbc-products/all.products.json", "去重后的完整 JSON 数据集"),
        ("抓取规模", "138 页 / 1,375 条", "官网报告、解析记录、唯一代码均为 1,375；零重复、零失败页"),
        ("发行方结构", "代销工银理财 776 条；其他具名发行方 598 条；发行方空白 1 条", "公开销售货架，不是纯自营产品主数据"),
        ("销售状态", "开放中 1,196 条；募集中 39 条；其余为最近开放日等状态", "状态是快照时点信息"),
        ("风险标签", "PR1 318 条；PR2 865 条；PR3 192 条；PR4/PR5 0 条", "整体偏稳健，不能外推高风险覆盖"),
    ]
    add_matrix_table(doc, ["指标", "结果", "解释"], profile_rows, [1800, 2800, 4760], font_size=9.0)
    add_para(doc, "风险分布：PR1 约 23.1%，PR2 约 62.9%，PR3 约 14.0%。这说明当前快照更适合验证稳健型到平衡型理财货架，"
             "不适合单独用于证明塔尖权益、另类或高风险交易产品的广度。", size=10.3, color=NAVY, bold=True)

    add_heading(doc, "四、工行可以做到的产品类型", 1)
    product_rows = [
        ("现金管理", "名称含“现金管理” 71 条；另有大量日开、活钱、添益类记录", "可做日开/短期流动性产品筛选、期限和起购金额比较", "高"),
        ("固定收益", "名称含“固定收益”或“固收” 834 条；明确“固定收益类” 154 条", "可作为稳健配置主力，覆盖多期限、开放/封闭及 PR1-PR3", "高"),
        ("混合与权益增强", "“混合” 80 条；“权益/指数/红利”相关名称 42 条", "可做混合、红利/指数增强等策略层比较；不等于直接股票交易", "中"),
        ("FOF / ESG", "FOF 21 条；ESG 43 条", "可作为多资产或主题策略补充；仍需说明书确认底层资产和比例", "中"),
        ("QDII / 外币 / 全球配置", "QDII 8 条；美元、港币、欧元、澳元 24 条；广义跨境/全球/外币/QDII 49 条", "可做有限跨境暴露和币种筛选；不等于完整全球资产平台", "中低"),
        ("指数挂钩 / 对冲", "指数挂钩 2 条；名称含“对冲” 21 条；直接“衍生/期权/期货/掉期” 0 条", "可识别封装策略名称；不能宣称客户直接开展衍生品交易", "低"),
        ("全权委托", "仅 1 条 PBZT0773；发行方空白，最近开放日早于快照日，起购 2,700 万元", "只能作为待核验线索，不足以支撑当前可办或自动推荐", "极低"),
        ("另类投资", "“另类/私募/股权/REIT/不动产/实物”等关键词均未检出", "本次清单无可验证证据", "未证实"),
    ]
    add_matrix_table(doc, ["产品类型", "清单证据", "可做到的动作/边界", "完善度"], product_rows, [1700, 3050, 3650, 960], font_size=8.4)
    add_para(doc, "注：关键词统计存在重叠，不能相加为独立产品数量；名称标签也不能替代产品分类、底层资产和适当性核验。", size=9.2, color=MUTED, italic=True, after=8)

    add_heading(doc, "五、与图片 13 类能力的逐项对比", 1)
    benchmark_rows = [
        ("股权投资", "名称中可见混合、权益/指数/红利、打新等策略描述；直接“权益类”名称 0 条", "可做权益增强或混合策略的产品筛选；不能宣称直接股票交易、全球市场接入或独立股权投资平台", "中低"),
        ("固定收益策略", "固定收益/固收关键词合计 834 条；风险和期限字段完整", "可以作为当前产品能力的核心层，支持现金管理到中长期稳健配置", "高"),
        ("衍生产品交易", "直接衍生/期权/期货/掉期 0 条；指数挂钩 2 条", "仅能识别少量封装或挂钩策略，不能证明面向客户的直接衍生品交易服务", "低"),
        ("全权委托投资组合", "仅 1 条全权委托记录，且状态和发行方存在核验问题", "只能作为人工核验线索；不能据此形成当前全权委托产品池", "极低"),
        ("另类投资", "另类、私募、REIT、不动产、实物等关键词均未检出", "本数据源不能支撑另类投资推荐", "未证实"),
        ("保险规划", "理财列表中保险关键词为 0，缺少保单和保障字段", "需要保险产品目录、保障责任、现金价值和适当性数据；当前清单不支持", "不在本数据源"),
        ("信托和遗产规划服务", "信托、遗产关键词为 0", "需要信托产品、法律结构、受益人和资产权属数据；当前清单不支持", "不在本数据源"),
        ("市场研究和分析", "列表没有研究报告、宏观分析或市场快讯字段", "可以做静态产品字段比较，但不能产出完整投研服务", "低"),
        ("机构客户服务", "来源明确为个人网银列表；机构/法人/企业客户关键词为 0", "机构客户能力需要企业金融或机构理财产品目录，当前清单不支持", "不在本数据源"),
        ("全面传承规划", "传承关键词为 0；无家庭成员、资产权属、受益安排字段", "需要家族信托、遗嘱/受益安排和法律服务协同，当前清单不支持", "不在本数据源"),
        ("人寿保险", "寿险、年金、重疾、医疗等保险关键词均为 0", "不能由理财列表证明人寿保险产品或服务能力", "不在本数据源"),
        ("税务责任规划", "税务、税收、节税等关键词均为 0", "需要税务居民身份、司法辖区和规则引擎，当前清单不支持", "不在本数据源"),
        ("财富结构设计", "无财富结构、持股架构、家族企业或税法字段", "当前清单可做投资产品配置，不能完成家庭资产负债、持股架构、信托和税务结构设计", "不在本数据源"),
    ]
    add_matrix_table(doc, ["图片类型", "可核对证据", "结论与边界", "完善度"], benchmark_rows, [1800, 2820, 3780, 960], font_size=8.2)

    add_heading(doc, "六、完善程度评价", 1)
    add_para(doc, "本报告采用四档评价：高 = 有较多记录、状态/风险/期限等核心字段可用，能够形成可复核的产品筛选；"
             "中 = 有多种产品或策略，但覆盖范围、复杂度或专业字段有限；低 = 只有少量名称线索或仅能做静态展示；"
             "未证实/不在本数据源 = 当前数据不能支持结论，需补充数据源或人工核验。")
    maturity_rows = [
        ("产品数量与销售货架", "高", "1,375 条唯一记录、1,196 条开放中；可按发行方、状态、期限、起购金额和风险做基础筛选。"),
        ("产品元数据可用性", "中高", "产品代码、名称、状态、主指标、起购金额、期限、风险、来源和抓取时间均为 100% 非空；发行方 99.9% 非空。"),
        ("收益与绩效比较", "低", "业绩比较基准仅 87 条（6.3%），单位净值 1 条（0.1%），每万份收益/年化收益各 173 条（12.6%）；不适合完整绩效分析。"),
        ("风险与策略宽度", "中等偏稳健", "PR1-PR3 全覆盖，但无 PR4/PR5；混合、指数增强和少量挂钩策略可见，直接权益/另类证据不足。"),
        ("跨境与币种", "中低", "有 QDII、全球配置和多币种名称，但数量有限，且缺少汇率、税务、地区适当性和交易规则字段。"),
        ("顾问服务闭环", "未证实", "缺少家庭资产负债、保险保障、信托、传承、税务、机构服务和投研服务字段。"),
    ]
    add_matrix_table(doc, ["评价维度", "评级", "依据"], maturity_rows, [2350, 1650, 5360], font_size=8.8)

    add_heading(doc, "七、基于当前数据可以落地的功能", 1)
    for text in [
        "产品货架检索：按风险标签、期限/最短持有期、起购金额、销售状态、发行方和币种进行筛选。",
        "稳健型配置：以现金管理和固定收益为底层，叠加混合、FOF、ESG或红利/指数增强策略形成分层组合。",
        "产品横向比较：展示主指标、期限、起购金额、开放安排和风险等级；对于缺失的历史净值、费率和业绩基准明确标注“暂无数据”。",
        "有限跨境展示：识别 QDII、美元、港币、欧元、澳元及全球配置产品，并把汇率和适当性核验设置为前置提示。",
        "人工转介：当用户提出保险、信托、传承、税务、机构服务等需求时，进入“待接入数据源/专业顾问转介”流程，不伪造自动推荐结果。",
    ]:
        add_bullet(doc, text)

    add_heading(doc, "八、当前数据不能直接支持的功能", 1)
    for text in [
        "不能把理财名称中的“权益”“指数”“对冲”解释为客户可直接交易股票、期权、期货或其他衍生品。",
        "不能仅凭本次列表完成家庭资产负债表、负债压力、保险保障缺口或多目标资金冲突的完整测算。",
        "不能生成信托/遗产/传承方案、税务责任规划或财富结构设计；这些需要法律、税务、信托和客户身份数据。",
        "不能把单条全权委托记录当作当前可办服务；PBZT0773 需要核对产品有效性、发行方、客户资格和交易状态。",
        "不能用当前列表字段做完整风险收益分析、回撤分析、历史业绩归因或组合自动再平衡。",
    ]:
        add_bullet(doc, text)

    add_heading(doc, "九、补齐产品与服务能力的建议", 1)
    roadmap_rows = [
        ("第一优先级：产品主数据", "补充标准化产品类别、底层资产、币种、费率、申赎规则、风险揭示书、说明书、历史净值和业绩基准。", "把“可展示”提升为“可比较、可适当性匹配”。"),
        ("第二优先级：保险与养老", "接入保险产品责任、保额、缴费、现金价值、退保损失、养老金和个人养老金产品字段。", "支持保障缺口和长期现金流规划，但仍需合规审核。"),
        ("第三优先级：信托/传承/税务", "建立信托、受益人、资产权属、税务居民身份、司法辖区和专业顾问转介接口。", "把图片中的专业服务从“未证实”变成“可核验的服务编排”。"),
        ("第四优先级：研究与机构", "接入市场资讯、宏观指标、投研观点、机构/企业金融目录和私银专属服务权限。", "支持市场研究、机构客户和高净值客户的分层服务。"),
        ("全流程控制", "对每项能力标注“已验证/间接证据/待接入/人工核验”，并保留来源、时间戳、适当性和消保提示。", "降低误推荐、过度宣传和把产品标签当作服务能力的风险。"),
    ]
    add_matrix_table(doc, ["阶段", "需要补充", "预期提升"], roadmap_rows, [2200, 4600, 2560], font_size=8.7)

    add_heading(doc, "十、最终判断", 1)
    add_para(doc, "如果问题是“工行现有公开理财产品能否支撑一个家庭财富管理工具的产品推荐层”，答案是：可以，且以现金管理和固定收益为核心的部分较成熟；"
             "混合、FOF、ESG、红利/指数增强和有限跨境产品可以作为补充。")
    add_para(doc, "如果问题是“工行是否已经在一套公开产品清单中完整覆盖图片所示的财富管理服务”，答案是：不能从本次爬取结果得出。"
             "保险规划、信托与传承、税务、财富结构、市场研究和机构客户服务应被建模为跨产品域的专业服务能力，并通过额外数据源、权限和人工核验完成闭环。", color=NAVY, bold=True)
    add_callout(doc, "建议对外表述", "“工行公开理财货架已形成以现金管理和固定收益为主、混合/权益增强与有限跨境策略为补充的产品供给；"
                "保险、信托、传承、税务、机构和投研服务需要在相应产品域和专业服务数据接入后再做能力承诺。”", fill="EAF4EF", accent=GREEN)

    doc.add_page_break()
    add_heading(doc, "附录 A：统计与字段质量", 1)
    field_rows = [
        ("核心识别与销售字段", "产品代码、名称、销售状态、主指标、起购金额、期限、风险等级、来源 URL、抓取时间", "均为 100% 非空"),
        ("发行方", "issuerLabel", "1,374 / 1,375 非空（99.9%）；PBZT0773 为空"),
        ("业绩比较基准", "benchmarkValue", "87 / 1,375（6.3%）"),
        ("单位净值", "unitNav", "1 / 1,375（0.1%）"),
        ("每万份收益 / 年化收益", "per10kIncome / annualizedYield", "各 173 / 1,375（12.6%）"),
        ("metrics 结构", "页面指标数组", "1,201 条有 3 个指标；174 条有 4 个指标"),
        ("异常记录", "PBZT0773；20HH6610", "前者发行方缺失；后者“最近开放日”字段缺日期"),
        ("缺失字段类型", "产品类别、底层资产、费率、持仓、历史净值、说明书、风险揭示书、交易资格等", "不适合直接完成投顾或合规推荐"),
    ]
    add_matrix_table(doc, ["字段组", "字段/记录", "质量结论"], field_rows, [2500, 3500, 3360], font_size=8.8)

    add_heading(doc, "附录 B：可复核证据记录", 1)
    evidence_rows = [
        ("现金管理", "20G2901B", "ICBC-WMP/data/icbc-products/pages/all/page-0001.json", "现金管理/开放型理财记录"),
        ("混合类", "20HH3151", "ICBC-WMP/data/icbc-products/pages/all/page-0127.json", "混合类、定期开放"),
        ("QDII 美元", "230701", "ICBC-WMP/data/icbc-products/pages/all/page-0089.json", "高盛工银理财盛景 QDII 私银美元产品"),
        ("ESG 固收", "25G2097A", "ICBC-WMP/data/icbc-products/pages/all/page-0040.json", "ESG 相关固定收益产品"),
        ("多资产 FOF", "19GS5881", "ICBC-WMP/data/icbc-products/pages/all/page-0070.json", "多资产 FOF 固收类产品"),
        ("指数挂钩", "HH1904 / 20GS6150", "ICBC-WMP/data/icbc-products/pages/all/page-0126.json；page-0128.json", "指数挂钩或股债轮动策略"),
        ("全权委托线索", "PBZT0773", "ICBC-WMP/data/icbc-products/pages/all/page-0125.json", "私银客户专属全权委托；需人工核验"),
    ]
    add_matrix_table(doc, ["能力线索", "产品代码", "文件位置", "用途"], evidence_rows, [1700, 1550, 3750, 2360], font_size=8.2)

    add_heading(doc, "附录 C：来源与使用说明", 1)
    add_source_list(doc, [
        ("数据范围说明", "ICBC-WMP/data/icbc-products/README.md", "说明来源、口径、字段和不可外推边界"),
        ("抓取统计", "ICBC-WMP/data/icbc-products/all.manifest.json", "138 页、1,375 条、零重复、零失败"),
        ("校验结果", "ICBC-WMP/data/icbc-products/all.validation.json", "valid=true，errors/warnings 为空"),
        ("完整产品记录", "ICBC-WMP/data/icbc-products/all.products.json", "本报告统计与示例记录的主要证据"),
        ("项目内产品体系调研", "金融产品/11_原始模块_工商银行个人金融产品体系_20260813.md", "更广产品体系的背景资料；未与本次爬取清单混为同一证据口径"),
        ("项目内财富工具调研", "金融产品/12_原始模块_工行财富管理工具调研_20260813.md", "服务和工具背景资料；不等于本次产品爬取结果"),
    ])
    add_para(doc, "免责声明：本报告用于项目研究和产品能力边界分析，不构成投资建议、销售话术或对具体产品的推荐。"
             "产品状态、风险等级、收益指标和销售资格应以工行及相关发行机构最新官方披露为准。", size=9.2, color=MUTED, italic=True, before=8, after=0)

    doc.core_properties.title = "工商银行财富管理产品能力对标报告"
    doc.core_properties.subject = "工行公开理财产品与图片标杆财富管理服务类型对比"
    doc.core_properties.author = "项目研究组"
    doc.core_properties.keywords = "工商银行;财富管理;理财产品;产品能力;对标分析"
    doc.save(OUT_PATH)
    print(OUT_PATH)


if __name__ == "__main__":
    build_report()
